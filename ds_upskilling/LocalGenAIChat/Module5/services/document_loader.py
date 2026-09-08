from pathlib import Path
from pypdf import PdfReader
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation


class DocumentLoader:

    def __init__(self, document_folder):
        self.document_folder = Path(document_folder)

    def load_documents(self):

        documents = []

        for file in self.document_folder.iterdir():

            if not file.is_file():
                continue

            extension = file.suffix.lower()

            try:

                if extension == ".pdf":
                    content = self.load_pdf(file)

                elif extension == ".docx":
                    content = self.load_docx(file)

                elif extension == ".txt":
                    content = self.load_text(file)

                elif extension == ".md":
                    content = self.load_text(file)

                elif extension == ".xlsx":
                    content = self.load_excel(file)

                elif extension == ".pptx":
                    content = self.load_ppt(file)

                else:
                    print(f"Skipping unsupported file: {file.name}")
                    continue

                documents.append(
                    {
                        "file_name": file.name,
                        "content": content
                    }
                )

            except Exception as ex:

                print(f"Error reading {file.name}: {ex}")

        return documents

    def load_pdf(self, file):

        reader = PdfReader(file)

        text = ""

        for page in reader.pages:

            page_text = page.extract_text()

            if page_text:
                text += page_text + "\n"

        return text

    def load_docx(self, file):

        document = Document(file)

        text = ""

        for paragraph in document.paragraphs:
            text += paragraph.text + "\n"

        return text

    def load_text(self, file):

        return file.read_text(
            encoding="utf-8",
            errors="ignore"
        )

    def load_excel(self, file):

        workbook = load_workbook(
            filename=file,
            data_only=True
        )

        text = ""

        for sheet in workbook.worksheets:

            text += f"\nSheet: {sheet.title}\n"

            for row in sheet.iter_rows(values_only=True):

                values = [
                    str(cell)
                    for cell in row
                    if cell is not None
                ]

                text += " | ".join(values)

                text += "\n"

        return text

    def load_ppt(self, file):

        presentation = Presentation(file)

        text = ""

        for slide in presentation.slides:

            for shape in slide.shapes:

                if hasattr(shape, "text"):

                    text += shape.text

                    text += "\n"

        return text